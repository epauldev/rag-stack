"""Génère des documents d'exemple (PDF, DOCX, EPUB, TXT, PDF scanné) pour les tests."""
from __future__ import annotations

from pathlib import Path

BIO_P1 = (
    "La photosynthèse est le processus par lequel les plantes vertes convertissent l'énergie lumineuse en énergie "
    "chimique. Elle se déroule dans les chloroplastes, des organites riches en chlorophylle. La phase claire "
    "produit de l'ATP et du NADPH, tandis que le cycle de Calvin fixe le dioxyde de carbone dans le stroma. "
    "Les pigments absorbent surtout la lumière rouge et bleue et réfléchissent le vert, ce qui explique la "
    "couleur des feuilles. Sans lumière, la phase claire s'arrête et le cycle de Calvin ralentit rapidement. "
)
BIO_P2 = (
    "La respiration cellulaire se déroule dans les mitochondries. Le glucose est dégradé par la glycolyse, puis "
    "par le cycle de Krebs, et la chaîne respiratoire produit l'essentiel de l'ATP. L'oxygène est l'accepteur "
    "final d'électrons. Les levures peuvent aussi réaliser une fermentation alcoolique lorsque l'oxygène manque. "
)
DB_TEXT = [
    ("Vector databases", 1),
    ("A vector database stores embeddings and answers nearest neighbor queries. The HNSW index builds a "
     "hierarchical navigable small world graph that gives fast approximate nearest neighbor search. Recall can "
     "be tuned with the ef_search parameter at the cost of latency. Product quantization compresses vectors to "
     "reduce memory. " * 3, 0),
    ("Hybrid retrieval", 1),
    ("Hybrid retrieval combines dense vectors with BM25 lexical search. Reciprocal rank fusion merges the two "
     "rankings without calibrating scores. A cross-encoder reranker then reorders the best candidates for "
     "precision. Exact identifiers such as the error code ZX-4471 are found by lexical search, not by embeddings. " * 3, 0),
]
ECO_HTML = """<html><body>
<h1>Politique monétaire</h1>
<p>L'inflation désigne la hausse générale et durable des prix. La banque centrale relève ses taux directeurs
pour freiner la demande et ramener l'inflation vers sa cible de deux pour cent. Les anticipations des agents
jouent un rôle central dans la transmission de la politique monétaire.</p>
<p>Le taux de change et le canal du crédit complètent les mécanismes de transmission. Une politique trop
restrictive peut provoquer une récession, ce qui illustre l'arbitrage entre inflation et activité.</p>
<h1>Dette publique</h1>
<p>La soutenabilité de la dette publique dépend du différentiel entre le taux d'intérêt et la croissance.
Un excédent primaire est nécessaire lorsque le taux d'intérêt dépasse durablement la croissance.</p>
</body></html>"""
NOTES_TXT = (
    "Notes de lecture sans titres. Le théorème central limite explique la convergence des moyennes "
    "empiriques vers une loi normale. L'intervalle de confiance à 95 pour cent utilise le quantile 1,96. "
) * 4


def make_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 80), "Chapitre 1 La photosynthese", fontsize=22)
    page.insert_textbox(pymupdf.Rect(72, 100, 540, 760), BIO_P1 * 3, fontsize=11)
    page = doc.new_page()
    page.insert_text((72, 80), "Chapitre 2 La respiration cellulaire", fontsize=22)
    page.insert_textbox(pymupdf.Rect(72, 100, 540, 760), BIO_P2 * 6, fontsize=11)
    doc.save(path)
    doc.close()


def make_scanned_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    for _ in range(3):
        page = doc.new_page()
        page.draw_rect(pymupdf.Rect(100, 100, 400, 500), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
    doc.save(path)
    doc.close()


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    for text, level in DB_TEXT:
        if level:
            d.add_heading(text, level=level)
        else:
            d.add_paragraph(text)
    d.save(path)


def make_epub(path: Path) -> None:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("id-eco")
    book.set_title("Economie")
    book.set_language("fr")
    ch = epub.EpubHtml(title="Chapitre", file_name="ch1.xhtml", lang="fr")
    ch.content = ECO_HTML
    book.add_item(ch)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", ch]
    epub.write_epub(str(path), book)


def make_all(folder: Path) -> dict[str, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    files = {
        "pdf": folder / "Campbell_Biologie_2020.pdf",
        "docx": folder / "Kleppmann_Vector_Databases.docx",
        "epub": folder / "Mankiw_Economie.epub",
        "txt": folder / "Notes_statistiques.txt",
        "scan": folder / "Scan_sans_texte.pdf",
    }
    make_pdf(files["pdf"])
    make_docx(files["docx"])
    make_epub(files["epub"])
    files["txt"].write_text(NOTES_TXT, encoding="utf-8")
    make_scanned_pdf(files["scan"])
    return files
